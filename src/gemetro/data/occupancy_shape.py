"""I-06: hourly occupancy shape I(h)/I_pico per line (PHASE_0_REPORT.md, "Ocupación N"; L-02).

I(h)    = A(line, h) / [f_ida(h) + f_vuelta(h)]
I_pico  = max over h in 4..22 of I(h)

A: median over weekdays (Monday to Friday, holidays not excluded) of January to September
2025 of the hourly entries of the line (S-03). f: scheduled trains per hour and direction
in the GTFS "Laboral" service, counted by departure time at the counting station (S-15):
Parque Berrío (BER) for line A, San Antonio (SAB) for line B.

I(h) is a relative load index (entries per train departure), not passengers on board.
"""

from __future__ import annotations

import datetime as dt
import zipfile
from dataclasses import dataclass

import openpyxl
import pandas as pd

from gemetro.data.medellin_download import MEDELLIN_RAW
from gemetro.paths import ARTIFACTS

OCCUPANCY_ARTIFACT = ARTIFACTS / "occupancy_shape.csv"
AFLUENCIA_FILE = "afluencia_2025.xlsx"
GTFS_FILE = "gtfs_2025.zip"

HOURS = range(4, 23)  # scenario hours 4..22 (PHASE_1_ARCHITECTURE.md §9.2)
MONTHS = range(1, 10)  # January to September 2025


@dataclass(frozen=True)
class Line:
    name: str  # scenario name: "A" or "B"
    afluencia_label: str
    gtfs_route: str
    counting_stop: str


LINES = (
    Line("A", "LÍNEA A", "LA", "BER"),
    Line("B", "LÍNEA B", "LB", "SAB"),
)


class OccupancyShapeError(ValueError):
    """Raised when the data cannot produce a valid shape (stop point, §9 of the plan)."""


def read_afluencia(path) -> pd.DataFrame:
    """Long table (date, line, hour, entries) from the afluencia workbook.

    The workbook has a two-row header: row 1 names, row 2 the hour of each column.
    """
    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    hour_row = rows[1]
    hour_columns = {i: v.hour for i, v in enumerate(hour_row) if isinstance(v, dt.time)}
    records = []
    for row in rows[2:]:
        if not isinstance(row[0], dt.datetime):
            continue
        for column, hour in hour_columns.items():
            records.append((row[0].date(), row[1], hour, row[column]))
    return pd.DataFrame(records, columns=["date", "line", "hour", "entries"])


def weekday_median_entries(afluencia: pd.DataFrame, label: str) -> pd.Series:
    data = afluencia[afluencia["line"] == label]
    dates = pd.to_datetime(data["date"])
    keep = (dates.dt.dayofweek < 5) & dates.dt.month.isin(MONTHS) & (dates.dt.year == 2025)
    data = data[keep].copy()
    data["entries"] = pd.to_numeric(data["entries"]).fillna(0)
    return data.groupby("hour")["entries"].median()


def trains_per_hour(gtfs_zip, line: Line) -> pd.DataFrame:
    """Scheduled departures per hour and direction at the counting stop, service Laboral."""
    with zipfile.ZipFile(gtfs_zip) as archive:
        names = {name.rsplit("/", 1)[-1]: name for name in archive.namelist()}
        trips = pd.read_csv(archive.open(names["trips.txt"]), dtype=str)
        stop_times = pd.read_csv(archive.open(names["stop_times.txt"]), dtype=str)
    trips = trips[(trips["route_id"] == line.gtfs_route) & (trips["service_id"] == "Laboral")]
    stops = stop_times[stop_times["stop_id"] == line.counting_stop]
    stops = stops.merge(trips[["trip_id", "direction_id"]], on="trip_id")
    stops["hour"] = stops["departure_time"].str.split(":").str[0].astype(int)
    counts = stops.groupby(["hour", "direction_id"]).size().unstack(fill_value=0)
    return counts.reindex(columns=["0", "1"], fill_value=0)


def line_shape(entries: pd.Series, trains: pd.DataFrame, line: Line) -> pd.DataFrame:
    table = pd.DataFrame(index=pd.Index(list(HOURS), name="hour"))
    table["entries"] = entries.reindex(table.index)
    table["trains_dir0"] = trains["0"].reindex(table.index).fillna(0).astype(int)
    table["trains_dir1"] = trains["1"].reindex(table.index).fillna(0).astype(int)
    missing_entries = table.index[table["entries"].isna() | (table["entries"] <= 0)].tolist()
    no_trains = table.index[(table["trains_dir0"] == 0) | (table["trains_dir1"] == 0)].tolist()
    if missing_entries or no_trains:
        raise OccupancyShapeError(
            f"line {line.name}: hours without entries {missing_entries}, "
            f"hours without trains in some direction {no_trains}"
        )
    table["load_index"] = table["entries"] / (table["trains_dir0"] + table["trains_dir1"])
    table["load_ratio"] = table["load_index"] / table["load_index"].max()
    table.insert(0, "line", line.name)
    return table.reset_index()


def build_occupancy_shape() -> pd.DataFrame:
    afluencia = read_afluencia(MEDELLIN_RAW / AFLUENCIA_FILE)
    tables = []
    for line in LINES:
        entries = weekday_median_entries(afluencia, line.afluencia_label)
        trains = trains_per_hour(MEDELLIN_RAW / GTFS_FILE, line)
        tables.append(line_shape(entries, trains, line))
    return pd.concat(tables, ignore_index=True)


def main() -> None:
    table = build_occupancy_shape()
    OCCUPANCY_ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(OCCUPANCY_ARTIFACT, index=False, float_format="%.6f", lineterminator="\n")
    print(f"wrote {OCCUPANCY_ARTIFACT} ({len(table)} rows)")


if __name__ == "__main__":
    main()
